"""Leitura e gravação da planilha Excel (inscrições e ganhadores)."""

import re
import socket
import sys
import time
import unicodedata
import zipfile
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import from_excel
from openpyxl.utils.exceptions import InvalidFileException

# ---------------------------------------------------------------------------
# Nome do arquivo Excel (gerado pelo Microsoft Forms)
# ---------------------------------------------------------------------------
NOME_ARQUIVO = "Presença Palestras, Atividades e Oficinas SIPAT 2026.xlsx"

# ---------------------------------------------------------------------------
# Índices dos campos nas tuplas devolvidas por ler_inscricoes.
# A posição de cada coluna na planilha do Forms é descoberta pelo cabeçalho
# (ver _mapear_colunas), então colunas extras ou fora de ordem não quebram
# a leitura. Estes índices são os da tupla já montada, não os da planilha.
# ---------------------------------------------------------------------------
COL_HORA = 0        # hora de conclusão (momento do envio), datetime ou None
COL_NOME = 1
COL_EMPRESA = 2
COL_MATRICULA = 3
COL_EVENTO = 4

# Palavras procuradas nos cabeçalhos (comparação sem acento e minúsculas).
# Cada campo pega a primeira coluna ainda livre que contenha uma das palavras,
# exceto "nome", que pega a última: o Forms tem uma coluna automática "Nome"
# (conta de quem respondeu) antes das perguntas do formulário.
CABECALHOS_ESPERADOS = [
    ("hora",      ["hora de conclusao", "completion time"]),
    ("matricula", ["matricula", "chave"]),
    ("empresa",   ["empresa"]),
    ("evento",    ["evento", "palestra", "atividade", "oficina"]),
    ("nome",      ["nome", "name"]),
]
CAMPOS_OBRIGATORIOS = {"hora", "matricula", "evento", "nome"}
DESCRICAO_CAMPOS = {
    "hora": "Hora de conclusão",
    "matricula": "Matrícula",
    "empresa": "Empresa",
    "evento": "Evento / Palestra / Atividade / Oficina",
    "nome": "Nome",
}

# Formatos aceitos quando a hora vem como texto em vez de data do Excel
FORMATOS_HORA = [
    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M",
    "%m/%d/%y %H:%M:%S", "%m/%d/%y %H:%M",
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
]

# ---------------------------------------------------------------------------
# Ganhadores
# Cada sessão encerrada gera um arquivo próprio na pasta "Ganhadores/" ao lado
# da planilha (sincronizada pelo OneDrive). A planilha do Forms nunca é
# reescrita, então vários computadores podem gravar ao mesmo tempo sem
# conflito de sincronização.
# ---------------------------------------------------------------------------
CABECALHO_GANHADORES = ["Evento", "Nome", "Matrícula", "Empresa", "Prêmio", "Assinatura"]

ABA_GANHADORES = "Ganhadores"
PASTA_GANHADORES = "Ganhadores"

# Tentativas de abrir um arquivo bloqueado (ex.: OneDrive sincronizando)
TENTATIVAS_LEITURA = 6
ESPERA_ENTRE_TENTATIVAS = 1.0  # segundos

# Início de um arquivo criptografado (senha ou rótulo de confidencialidade):
# o .xlsx vira um contêiner OLE em vez de um ZIP.
ASSINATURA_OLE = b"\xd0\xcf\x11\xe0"


# ---------------------------------------------------------------------------
# Texto
# ---------------------------------------------------------------------------
def normalizar_texto(valor) -> str:
    """Texto para comparação: sem acentos, minúsculo, espaços simples (RN-04)."""
    if valor is None:
        return ""
    texto = unicodedata.normalize("NFKD", str(valor))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return " ".join(texto.casefold().split())


def eventos_da_celula(valor) -> list:
    """Eventos de uma célula. Em perguntas de múltipla seleção o Forms grava
    as opções marcadas separadas por ";" (ex.: "Evento A;Evento B;")."""
    if valor is None:
        return []
    return [parte.strip() for parte in str(valor).split(";") if parte.strip()]


# ---------------------------------------------------------------------------
# Caminho da planilha
# ---------------------------------------------------------------------------
def caminho_esperado() -> Path:
    """Caminho esperado da planilha (sem abrir seletor)."""
    if getattr(sys, "frozen", False):
        pasta = Path(sys.executable).parent
    else:
        pasta = Path(
            r"C:\Users\FY0E\PETROBRAS"
            r"\Anna Clara Gomes Silva Sousa - PrestServ - SIPAT 2026"
        )
    return pasta / NOME_ARQUIVO


def caminho_da_planilha() -> Path:
    """Retorna o Path do arquivo Excel.

    Tenta o caminho esperado; se não existir, abre seletor de arquivos.
    """
    caminho = caminho_esperado()
    if caminho.exists():
        return caminho

    from tkinter import filedialog
    arquivo = filedialog.askopenfilename(
        title="Escolher planilha",
        filetypes=[("Excel", "*.xlsx")],
    )
    if arquivo:
        return Path(arquivo)

    return caminho


def pasta_ganhadores(arquivo_excel) -> Path:
    """Pasta onde ficam os arquivos de ganhadores (ao lado da planilha)."""
    return Path(arquivo_excel).parent / PASTA_GANHADORES


# ---------------------------------------------------------------------------
# Abertura com nova tentativa
# ---------------------------------------------------------------------------
def _abrir_para_leitura(arquivo_excel):
    """Abre o Excel em modo leitura, tentando de novo se estiver bloqueado.

    O arquivo é copiado inteiro para a memória e fechado na hora: assim o
    sistema não segura o arquivo aberto enquanto o OneDrive tenta atualizá-lo.
    Enquanto o OneDrive sincroniza, a leitura pode falhar (arquivo bloqueado)
    ou vir pela metade (ZIP inválido) — nos dois casos tenta de novo.
    """
    arquivo_excel = Path(arquivo_excel)
    ultimo_erro = None

    for tentativa in range(TENTATIVAS_LEITURA):
        try:
            dados = arquivo_excel.read_bytes()
        except FileNotFoundError:
            raise FileNotFoundError(
                f"Arquivo não encontrado:\n{arquivo_excel}\n\n"
                "Verifique se a planilha está na mesma pasta do programa."
            )
        except OSError as e:
            ultimo_erro = e
        else:
            if dados.startswith(ASSINATURA_OLE):
                raise PermissionError(
                    "A planilha está protegida e não pode ser lida:\n"
                    f"{arquivo_excel.name}\n\n"
                    "O arquivo tem senha ou um rótulo de confidencialidade "
                    "com criptografia. Remova a proteção no Excel "
                    "(Arquivo > Informações) e tente novamente."
                )
            try:
                return load_workbook(BytesIO(dados), read_only=True, data_only=True)
            except (zipfile.BadZipFile, InvalidFileException, KeyError) as e:
                ultimo_erro = e

        if tentativa < TENTATIVAS_LEITURA - 1:
            time.sleep(ESPERA_ENTRE_TENTATIVAS)

    raise PermissionError(
        f"Não foi possível abrir o arquivo:\n{arquivo_excel}\n\n"
        "Verifique se:\n"
        "• o arquivo não está aberto no Excel neste computador;\n"
        "• o OneDrive terminou de sincronizar (ícone sem setas azuis);\n"
        "• o antivírus/política da empresa não está bloqueando o programa.\n\n"
        "Depois tente novamente.\n\n"
        f"Detalhe técnico: {type(ultimo_erro).__name__}: {ultimo_erro}"
    )


# ---------------------------------------------------------------------------
# Leitura de inscrições
# ---------------------------------------------------------------------------
def _mapear_colunas(cabecalho) -> dict:
    """Descobre a posição de cada campo pelo texto do cabeçalho.

    Levanta ValueError listando os cabeçalhos encontrados se faltar
    algum campo obrigatório.
    """
    textos = [normalizar_texto(c) for c in cabecalho]
    usadas = set()
    mapa = {}

    for campo, palavras in CABECALHOS_ESPERADOS:
        candidatas = [
            i for i, texto in enumerate(textos)
            if i not in usadas and any(p in texto for p in palavras)
        ]
        if not candidatas:
            continue
        indice = candidatas[-1] if campo == "nome" else candidatas[0]
        mapa[campo] = indice
        usadas.add(indice)

    faltando = [c for c, _ in CABECALHOS_ESPERADOS
                if c in CAMPOS_OBRIGATORIOS and c not in mapa]
    if faltando:
        encontrados = "\n".join(
            f"• {c}" for c in cabecalho if c is not None
        ) or "(nenhum)"
        raise ValueError(
            "Não encontrei na primeira aba da planilha a(s) coluna(s):\n"
            + "\n".join(f"• {DESCRICAO_CAMPOS[c]}" for c in faltando)
            + f"\n\nCabeçalhos encontrados (linha 1):\n{encontrados}"
        )
    return mapa


def _para_datetime(valor):
    """Converte a célula de hora em datetime; None se vazia ou ilegível."""
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, (int, float)):
        try:
            return from_excel(valor)
        except (ValueError, OverflowError):
            return None
    if isinstance(valor, str):
        texto = valor.strip()
        for formato in FORMATOS_HORA:
            try:
                return datetime.strptime(texto, formato)
            except ValueError:
                pass
    return None


def ler_inscricoes(arquivo_excel):
    """Lê a primeira aba da planilha e devolve lista de tuplas.

    Cada tupla tem os campos nas posições COL_* (hora, nome, empresa,
    matrícula, evento), localizados pelo cabeçalho da linha 1. A hora vem
    como datetime (ou None). Linhas totalmente vazias são ignoradas.
    Os filtros de evento e horário ficam em sorteador.filtrar_inscricoes.
    """
    wb = _abrir_para_leitura(arquivo_excel)
    try:
        ws = wb.worksheets[0]  # sempre a primeira aba (nunca active)
        linhas = ws.iter_rows(values_only=True)

        cabecalho = next(linhas, None) or ()
        mapa = _mapear_colunas(cabecalho)

        def campo(valores, nome):
            indice = mapa.get(nome)
            if indice is None or indice >= len(valores):
                return None
            return valores[indice]

        inscricoes = []
        for valores in linhas:
            # Ignora linhas totalmente vazias
            if all(v is None for v in valores):
                continue

            inscricao = [None] * 5
            inscricao[COL_HORA] = _para_datetime(campo(valores, "hora"))
            inscricao[COL_NOME] = campo(valores, "nome")
            inscricao[COL_EMPRESA] = campo(valores, "empresa")
            inscricao[COL_MATRICULA] = campo(valores, "matricula")
            inscricao[COL_EVENTO] = campo(valores, "evento")
            inscricoes.append(tuple(inscricao))
    finally:
        wb.close()

    return inscricoes


# ---------------------------------------------------------------------------
# Leitura de ganhadores
# ---------------------------------------------------------------------------
def _linhas_da_aba(ws):
    """Linhas da aba a partir da linha 2, ignorando as totalmente vazias."""
    linhas = []
    primeira_linha = True
    for row in ws.iter_rows(min_row=1):
        if primeira_linha:
            primeira_linha = False
            continue

        valores = tuple(cell.value for cell in row)
        if all(v is None for v in valores):
            continue

        linhas.append(valores)
    return linhas


def ler_ganhadores(arquivo_excel):
    """Lê todos os ganhadores já gravados e devolve lista de tuplas.

    Junta:
      - a aba 'Ganhadores' da planilha principal, se existir (versões antigas
        do sistema gravavam lá);
      - a primeira aba de cada arquivo .xlsx da pasta 'Ganhadores/'.
    """
    ganhadores = []

    wb = _abrir_para_leitura(arquivo_excel)
    if ABA_GANHADORES in wb.sheetnames:
        ganhadores.extend(_linhas_da_aba(wb[ABA_GANHADORES]))
    wb.close()

    pasta = pasta_ganhadores(arquivo_excel)
    if pasta.is_dir():
        for arquivo in sorted(pasta.glob("*.xlsx")):
            # "~$..." são arquivos temporários do Excel (arquivo aberto)
            if arquivo.name.startswith("~$"):
                continue
            wb = _abrir_para_leitura(arquivo)
            ganhadores.extend(_linhas_da_aba(wb.worksheets[0]))
            wb.close()

    return ganhadores


# ---------------------------------------------------------------------------
# Gravação de ganhadores
# ---------------------------------------------------------------------------
def _destino_sessao(pasta, evento) -> Path:
    """Arquivo novo da sessão: evento + data/hora + computador.

    Se já existir um com o mesmo nome (duas sessões no mesmo segundo),
    acrescenta " (2)", " (3)"… — nunca sobrescreve ganhadores gravados.
    """
    evento_limpo = re.sub(r'[<>:"/\\|?*]', "-", str(evento)).strip(" .")
    agora = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    computador = socket.gethostname()
    base = f"{evento_limpo} - {agora} - {computador}"

    destino = pasta / f"{base}.xlsx"
    n = 2
    while destino.exists():
        destino = pasta / f"{base} ({n}).xlsx"
        n += 1
    return destino


def gravar_ganhadores(arquivo_excel, ganhadores):
    """Grava os ganhadores da sessão num arquivo novo em 'Ganhadores/'.

    - Cada ganhador é uma tupla (evento, nome, matricula, empresa, premio).
    - Adiciona coluna "Assinatura" vazia ao final de cada linha.
    - Nunca altera a planilha do Forms nem arquivos de outras sessões.

    Retorna o Path do arquivo criado.
    Levanta PermissionError com mensagem amigável se não conseguir gravar.
    """
    pasta = pasta_ganhadores(arquivo_excel)

    wb = Workbook()
    ws = wb.active
    ws.title = ABA_GANHADORES
    ws.append(CABECALHO_GANHADORES)
    for ganhador in ganhadores:
        ws.append(list(ganhador) + [""])

    try:
        pasta.mkdir(exist_ok=True)
        destino = _destino_sessao(pasta, ganhadores[0][0])
        wb.save(destino)
    except PermissionError:
        raise PermissionError(
            "Não foi possível salvar os ganhadores na pasta:\n"
            f"{pasta}\n\n"
            "Verifique se você tem permissão de edição na pasta "
            "compartilhada e tente novamente."
        )
    finally:
        wb.close()

    return destino
